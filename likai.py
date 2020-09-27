from pdfminer.pdfparser import PDFParser,PDFDocument
from pdfminer.pdfinterp import PDFResourceManager, PDFPageInterpreter
from pdfminer.converter import PDFPageAggregator
from pdfminer.layout import LTTextBoxHorizontal, LAParams
import os
import PyPDF2
from PyPDF2 import PdfFileReader,PdfFileWriter
import openpyxl as opl
import logging
import time
logging.basicConfig(level=logging.ERROR)



def mkResDir(path):
    savepath = os.path.join(path, '结果')
    # savepath = os.getcwd() + '结果'
    if os.path.exists(savepath):
        for file in os.listdir(savepath):
            os.remove(os.path.join(savepath,file))
    else:
        os.mkdir(savepath)

def main(path):
    files = os.listdir(path)
    # print(files)
    dic={}
    for file in files:
        if file.lower().endswith('.pdf'):
            L=[]
            path_file=os.path.join(path,file)
            print('当前处理=',path_file)
            ##########################提取学生信息部分--start##################
            print('*'*30)
            print('解析pdf开始')
            parser=PDFParser(open(path_file,'rb'))
            doc=PDFDocument()
            parser.set_document(doc)
            doc.set_parser(parser)
            doc.initialize()
            if doc.is_extractable:
                doc_resource=PDFResourceManager()
                doc_device=LAParams()
                doc_resource_device=PDFPageAggregator(doc_resource,laparams=doc_device)
                doc_interpreter=PDFPageInterpreter(doc_resource,doc_resource_device)
                for page in doc.get_pages():
                    # result=''
                    print('exec page')
                    doc_interpreter.process_page(page)
                    layout=doc_resource_device.get_result()
                    for x in layout:
                        print(type(x))
                        if isinstance(x,LTTextBoxHorizontal):
                            result=x.get_text().replace('\n','')
                            print(result)
                            if result.find('学号')>=0 and result.find('姓名')>=0:
                                xh=result.split('学号')[-1].split('姓名')[0]
                                xm=result.split('姓名')[-1].split('性别')[0]
                                L.append(xh+'#'+xm)
                        else:
                            print("x is not LTTextBox")
            else:
                print(file,'is Error!')
            parser.close()
            #########################提取学生信息部分--end############################
            #########################生成学生页码信息部分---start######################
            for index in range(len(L)):
                if L[index] not in dic.keys():
                    dic[L[index]]=str(L.index(L[index]))+'-'+str(index+L.count(L[index])-1)
            ########################处理学生页码信息部分----end#########################
            print('解析pdf结束。')
            print('拆分pdf开始！')
            ########################拆分pdf文件--start################################
            savepath = os.path.join(path, '结果')
            try:
                doc=PdfFileReader(open(path_file,'rb'))
                for k,v in dic.items():
                    pdf=PdfFileWriter()
                    start_page,end_page=int(v.split('-')[0]),int((v.split('-')[-1]))
                    for index in range(start_page,end_page+1):
                        page=doc.getPage(index)
                        pdf.addPage(page)
                    if os.path.exists(os.path.join(savepath,k.replace('#',' ')+'.pdf')):
                        os.remove(os.path.join(savepath,k.replace('#',' ')+'.pdf'))
                    with open(os.path.join(savepath,k.replace('#',' ')+'.pdf'),'wb') as f:
                        pdf.write(f)
                    f.close()
                print('拆分pdf结束！')
            except Exception as e:
                print('拆分pdf文件=',path_file,'失败!')
                print(e)

            # print(dic)
            ##################拆分pdf文件--end#########################################
            ##################生成拆分结果清单--开始############################################
            # content=[]
            print('*'*30)
            print('生成拆分结果清单开始!')
            try:
                if os.path.exists(os.path.join(path,'拆分结果清单.xlsx')):
                    os.remove(os.path.join(path,'拆分结果清单.xlsx'))
                wb=opl.Workbook()
                ws=wb.create_sheet('Res')
                ws.append(('学号','姓名','文件链接','收件人（自行录入）','方式（自行录入）'))
                for k,v in dic.items():
                    t=(
                        k.split('#')[0],
                        k.split('#')[-1],
                        '=hyperlink("'+os.path.join(savepath,k.replace('#',' ')+'.pdf')+'")',
                        '',
                    '')
                    ws.append(t)
                    # print(content)
                # ws.append(content)
                wb.save(os.path.join(os.getcwd(),'拆分结果清单.xlsx'))
            except Exception as e:
                print('生成拆分清单失败！请检查是否存在未关闭的“拆分结果清单.xlsx”文件！')
                print(e)
            print('生成拆分清单结束！')
            ##################

    return 1



if __name__=='__main__':
    print('开始！')
    path=os.getcwd()
    mkResDir(path)
    main(path)
    time.sleep(10)
    print('结束！')
